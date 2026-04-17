from __future__ import annotations

from io import BytesIO
from pathlib import Path
import sys
from tempfile import TemporaryDirectory
from unittest import TestCase

import duckdb
from openpyxl import Workbook


REPO_ROOT = Path(__file__).resolve().parents[1]
BDW_ROOT = REPO_ROOT / "bdw"
if str(BDW_ROOT) not in sys.path:
    sys.path.insert(0, str(BDW_ROOT))


from bit_data_workbench.backend.local_workspace_query_sources import (  # noqa: E402
    LocalWorkspaceQuerySourceManager,
    normalize_local_workspace_query_format,
)
from bit_data_workbench.config import Settings  # noqa: E402


def make_settings(database_path: Path) -> Settings:
    return Settings(
        service_name="bit-data-workbench",
        ui_title="DAAIFL Workbench",
        image_version="0.5.7",
        port=8000,
        duckdb_database=database_path,
        duckdb_extension_directory=database_path.parent / "duckdb-ext",
        max_result_rows=200,
        s3_endpoint="http://127.0.0.1:9000",
        s3_bucket="workspace",
        s3_access_key_id="minio",
        s3_access_key_id_file=None,
        s3_secret_access_key="miniosecret",
        s3_secret_access_key_file=None,
        s3_url_style="path",
        s3_use_ssl=False,
        s3_verify_ssl=False,
        s3_ca_cert_file=None,
        s3_session_token=None,
        s3_session_token_file=None,
        s3_startup_view_schema="s3",
        s3_startup_views=None,
        pg_host="127.0.0.1",
        pg_port="5432",
        pg_user="postgres",
        pg_password="postgres",
        pg_oltp_database="oltp",
        pg_olap_database="olap",
        pod_name=None,
        pod_namespace=None,
        pod_ip=None,
        node_name=None,
    )


class LocalWorkspaceQuerySourceManagerTests(TestCase):
    def test_sync_source_creates_queryable_csv_view_with_header_metadata(self) -> None:
        with TemporaryDirectory() as temp_dir:
            database_path = Path(temp_dir) / "workspace.duckdb"
            settings = make_settings(database_path)
            manager = LocalWorkspaceQuerySourceManager(settings=settings)
            connection = duckdb.connect(str(database_path))
            try:
                result = manager.sync_source(
                    conn=connection,
                    client_id="client-alpha",
                    entry_id="local-entry-1",
                    logical_relation="workspace.local.saved_results.local-entry-1",
                    file_name="tax-office.csv",
                    export_format="csv",
                    mime_type="text/csv",
                    file_bytes=(
                        b"record_id,canton_code,tax_office,assessed_amount_chf\n"
                        b"1,ZH,Zurich Central Tax Office,1200.50\n"
                        b"2,BE,Bern Regional Tax Office,918.25\n"
                    ),
                    csv_delimiter=",",
                    csv_has_header=True,
                )

                self.assertEqual(
                    [field.name for field in result.fields],
                    ["record_id", "canton_code", "tax_office", "assessed_amount_chf"],
                )
                self.assertEqual(
                    connection.execute(
                        f"SELECT record_id, canton_code, tax_office FROM {result.relation} ORDER BY record_id"
                    ).fetchall(),
                    [
                        (1, "ZH", "Zurich Central Tax Office"),
                        (2, "BE", "Bern Regional Tax Office"),
                    ],
                )
            finally:
                connection.close()

    def test_delete_and_clear_remove_registered_local_workspace_sources(self) -> None:
        with TemporaryDirectory() as temp_dir:
            database_path = Path(temp_dir) / "workspace.duckdb"
            settings = make_settings(database_path)
            manager = LocalWorkspaceQuerySourceManager(settings=settings)
            connection = duckdb.connect(str(database_path))
            try:
                first = manager.sync_source(
                    conn=connection,
                    client_id="client-alpha",
                    entry_id="local-entry-1",
                    logical_relation="workspace.local.saved_results.local-entry-1",
                    file_name="first.csv",
                    export_format="csv",
                    mime_type="text/csv",
                    file_bytes=b"id,name\n1,alpha\n",
                    csv_delimiter=",",
                    csv_has_header=True,
                )
                second = manager.sync_source(
                    conn=connection,
                    client_id="client-alpha",
                    entry_id="local-entry-2",
                    logical_relation="workspace.local.saved_results.local-entry-2",
                    file_name="second.csv",
                    export_format="csv",
                    mime_type="text/csv",
                    file_bytes=b"id,name\n2,beta\n",
                    csv_delimiter=",",
                    csv_has_header=True,
                )

                manager.delete_source(
                    conn=connection,
                    client_id="client-alpha",
                    entry_id="local-entry-1",
                )
                remaining_relations = connection.execute(
                    """
                    SELECT table_name
                    FROM information_schema.tables
                    WHERE table_schema = ?
                    ORDER BY table_name
                    """,
                    [second.relation.split(".")[0]],
                ).fetchall()
                self.assertEqual(remaining_relations, [(second.relation.split(".")[1],)])

                manager.clear_client_sources(
                    conn=connection,
                    client_id="client-alpha",
                )
                self.assertEqual(
                    connection.execute(
                        """
                        SELECT COUNT(*)
                        FROM information_schema.tables
                        WHERE table_schema = ?
                        """,
                        [first.relation.split(".")[0]],
                    ).fetchone(),
                    (0,),
                )
            finally:
                connection.close()

    def test_normalize_local_workspace_query_format_accepts_jsonl_as_json(self) -> None:
        self.assertEqual(
            normalize_local_workspace_query_format(
                file_name="saved-results.jsonl",
                export_format="",
                mime_type="application/json",
            ),
            "json",
        )

    def test_sync_source_creates_queryable_xml_view(self) -> None:
        with TemporaryDirectory() as temp_dir:
            database_path = Path(temp_dir) / "workspace.duckdb"
            settings = make_settings(database_path)
            manager = LocalWorkspaceQuerySourceManager(settings=settings)
            connection = duckdb.connect(str(database_path))
            try:
                result = manager.sync_source(
                    conn=connection,
                    client_id="client-alpha",
                    entry_id="local-entry-xml",
                    logical_relation="workspace.local.saved_results.local-entry-xml",
                    file_name="tax-office.xml",
                    export_format="xml",
                    mime_type="application/xml",
                    file_bytes=(
                        b"<rows>"
                        b"<row><record_id>1</record_id><tax_office>Zurich Central Tax Office</tax_office></row>"
                        b"<row><record_id>2</record_id><tax_office>Bern Regional Tax Office</tax_office></row>"
                        b"</rows>"
                    ),
                )

                self.assertEqual(
                    [field.name for field in result.fields],
                    ["record_id", "tax_office"],
                )
                self.assertEqual(
                    connection.execute(
                        f"SELECT record_id, tax_office FROM {result.relation} ORDER BY record_id"
                    ).fetchall(),
                    [
                        (1, "Zurich Central Tax Office"),
                        (2, "Bern Regional Tax Office"),
                    ],
                )
            finally:
                connection.close()

    def test_sync_source_creates_queryable_xlsx_view(self) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["record_id", "tax_office"])
        worksheet.append([1, "Zurich Central Tax Office"])
        worksheet.append([2, "Bern Regional Tax Office"])
        output = BytesIO()
        workbook.save(output)
        workbook.close()

        with TemporaryDirectory() as temp_dir:
            database_path = Path(temp_dir) / "workspace.duckdb"
            settings = make_settings(database_path)
            manager = LocalWorkspaceQuerySourceManager(settings=settings)
            connection = duckdb.connect(str(database_path))
            try:
                result = manager.sync_source(
                    conn=connection,
                    client_id="client-alpha",
                    entry_id="local-entry-xlsx",
                    logical_relation="workspace.local.saved_results.local-entry-xlsx",
                    file_name="tax-office.xlsx",
                    export_format="xlsx",
                    mime_type=(
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    ),
                    file_bytes=output.getvalue(),
                )

                self.assertEqual(
                    [field.name for field in result.fields],
                    ["record_id", "tax_office"],
                )
                self.assertEqual(
                    connection.execute(
                        f"SELECT record_id, tax_office FROM {result.relation} ORDER BY record_id"
                    ).fetchall(),
                    [
                        (1, "Zurich Central Tax Office"),
                        (2, "Bern Regional Tax Office"),
                    ],
                )
            finally:
                connection.close()
