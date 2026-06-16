from __future__ import annotations

from io import BytesIO
from pathlib import Path
import sys
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest import TestCase
from unittest.mock import patch

import duckdb
from openpyxl import Workbook


REPO_ROOT = Path(__file__).resolve().parents[1]
BDW_ROOT = REPO_ROOT / "bdw"
if str(BDW_ROOT) not in sys.path:
    sys.path.insert(0, str(BDW_ROOT))


from bit_data_workbench.backend.service import STARTUP_SEED_S3_BUCKETS, WorkbenchService  # noqa: E402
from bit_data_workbench.backend.s3_storage import s3_bucket_schema_name  # noqa: E402
from bit_data_workbench.backend.source_discovery import DiscoveredRelationSpec  # noqa: E402
from bit_data_workbench.backend.source_discovery import S3DataSourceDiscoverer  # noqa: E402
from bit_data_workbench.backend.source_discovery import build_s3_query  # noqa: E402
from bit_data_workbench.backend.source_discovery import drop_discovered_relation_object  # noqa: E402
from bit_data_workbench.config import Settings  # noqa: E402
from bit_data_workbench.models import SourceObject  # noqa: E402
from bit_data_workbench.version_info import current_repo_version  # noqa: E402


CURRENT_VERSION = current_repo_version(REPO_ROOT)


def make_settings() -> Settings:
    return Settings(
        service_name="bit-data-workbench",
        ui_title="DAAIFL Workbench",
        image_version=CURRENT_VERSION,
        port=8000,
        duckdb_database=Path("/tmp/workspace.duckdb"),
        duckdb_extension_directory=Path("/tmp/duckdb-ext"),
        service_consumption_data_dir=Path("/tmp/service-consumption"),
        service_consumption_cpu_memory_interval_seconds=3,
        service_consumption_s3_interval_seconds=3600,
        service_consumption_retention_hours=48,
        max_result_rows=200,
        s3_endpoint="http://127.0.0.1:9000",
        s3_bucket="vat-smoke-test",
        s3_access_key_id="minio",
        s3_access_key_id_file=None,
        s3_secret_access_key="miniosecret",
        s3_secret_access_key_file=None,
        s3_url_style="path",
        s3_use_ssl=False,
        s3_verify_ssl=False,
        s3_ca_cert_file=None,
        s3_session_token=None,
        s3_session_token_file=None,
        s3_startup_view_schema="s3",
        s3_startup_views=None,
        pg_host=None,
        pg_port=None,
        pg_user=None,
        pg_password=None,
        pg_oltp_database=None,
        pg_olap_database=None,
        pod_name=None,
        pod_namespace=None,
        pod_ip=None,
        node_name=None,
    )


def make_temp_settings(database_path: Path) -> Settings:
    settings = make_settings()
    settings.duckdb_database = database_path
    settings.duckdb_extension_directory = database_path.parent / "duckdb-ext"
    return settings


class FakeS3Client:
    def list_objects_v2(self, **kwargs):
        bucket = kwargs["Bucket"]
        prefix = str(kwargs.get("Prefix") or "")
        contents = []
        if bucket == "vat-smoke-test" and "incoming/tax-office.csv".startswith(prefix):
            contents.append({"Key": "incoming/tax-office.csv"})
        return {
            "Contents": contents,
            "IsTruncated": False,
        }

    def head_object(self, **kwargs):
        if kwargs["Bucket"] != "vat-smoke-test" or kwargs["Key"] != "incoming/tax-office.csv":
            raise AssertionError("Unexpected head_object request in test.")
        return {
            "ETag": '"abc123"',
            "ContentLength": 128,
            "Metadata": {
                "bdw_csv_has_header": "true",
                "bdw_csv_delimiter": "comma",
            }
        }


class DataExchangeHiddenS3Client:
    def list_objects_v2(self, **kwargs):
        bucket = kwargs["Bucket"]
        prefix = str(kwargs.get("Prefix") or "")
        contents = []
        if bucket == "vat-smoke-test":
            for key in (
                "--data-exchange--/files/secret/hidden.csv",
                "incoming/visible.csv",
            ):
                if key.startswith(prefix):
                    contents.append({"Key": key})
        return {
            "Contents": contents,
            "IsTruncated": False,
        }

    def head_object(self, **kwargs):
        if kwargs["Key"] == "incoming/visible.csv":
            return {
                "ETag": '"visible123"',
                "ContentLength": 32,
                "Metadata": {},
            }
        raise AssertionError("Hidden DataExchange object should not be inspected.")


class GeneratedMwaLoaderS3Client:
    keys = (
        "generated/mwa_abrechnung_test/csv/mwa_abrechnung_entities/part-00001.csv",
        "generated/mwa_abrechnung_test/csv/mwa_abrechnung_entities/part-00002.csv",
        "generated/mwa_abrechnung_test/json/mwa_abrechnung_entities/part-00001.jsonl",
        "generated/mwa_abrechnung_test/json/mwa_abrechnung_entities/part-00002.jsonl",
        "generated/mwa_abrechnung_test/parquet/mwa_abrechnung_entities/part-00001.parquet",
        "generated/mwa_abrechnung_test/parquet/mwa_abrechnung_entities/part-00002.parquet",
    )
    sizes = {
        "generated/mwa_abrechnung_test/csv/mwa_abrechnung_entities/part-00001.csv": 128,
        "generated/mwa_abrechnung_test/csv/mwa_abrechnung_entities/part-00002.csv": 96,
        "generated/mwa_abrechnung_test/json/mwa_abrechnung_entities/part-00001.jsonl": 256,
        "generated/mwa_abrechnung_test/json/mwa_abrechnung_entities/part-00002.jsonl": 192,
        "generated/mwa_abrechnung_test/parquet/mwa_abrechnung_entities/part-00001.parquet": 512,
        "generated/mwa_abrechnung_test/parquet/mwa_abrechnung_entities/part-00002.parquet": 384,
    }

    def list_objects_v2(self, **kwargs):
        bucket = kwargs["Bucket"]
        prefix = str(kwargs.get("Prefix") or "")
        if bucket != "vat-smoke-test":
            return {"Contents": [], "IsTruncated": False}
        return {
            "Contents": [
                {"Key": key}
                for key in self.keys
                if key.startswith(prefix)
            ],
            "IsTruncated": False,
        }

    def head_object(self, **kwargs):
        if kwargs["Bucket"] != "vat-smoke-test":
            raise AssertionError("Unexpected generated loader bucket.")
        key = kwargs["Key"]
        if key not in self.sizes:
            raise AssertionError(f"Unexpected generated loader head_object request: {key}")
        metadata = {}
        if key.endswith(".csv"):
            metadata = {
                "bdw_csv_has_header": "true",
                "bdw_csv_delimiter": "comma",
            }
        return {
            "ETag": f'"{Path(key).name}-etag"',
            "ContentLength": self.sizes[key],
            "Metadata": metadata,
        }


class SinglePartGeneratedMwaLoaderS3Client(GeneratedMwaLoaderS3Client):
    keys = (
        "generated/mwa_abrechnung_test/parquet/mwa_abrechnung_entities/part-00001.parquet",
    )
    sizes = {
        "generated/mwa_abrechnung_test/parquet/mwa_abrechnung_entities/part-00001.parquet": 512,
    }


class PartitionedParquetS3Client:
    keys = (
        "incoming/april/vat_smoke/tax_year=2025/data_0.parquet",
        "incoming/april/vat_smoke/tax_year=2026/data_0.parquet",
    )

    def list_objects_v2(self, **kwargs):
        bucket = kwargs["Bucket"]
        prefix = str(kwargs.get("Prefix") or "")
        if bucket != "vat-smoke-test":
            return {"Contents": [], "IsTruncated": False}
        return {
            "Contents": [
                {"Key": key}
                for key in self.keys
                if key.startswith(prefix)
            ],
            "IsTruncated": False,
        }

    def head_object(self, **kwargs):
        key = kwargs["Key"]
        if key not in self.keys:
            raise AssertionError(f"Unexpected partitioned parquet head_object request: {key}")
        return {
            "ETag": f'"{Path(key).parent.name}-etag"',
            "ContentLength": 256,
            "Metadata": {},
        }


class StartupSeedBucketS3Client:
    bucket = "poc-tests-performance-evaluation-mwa-abrechnung-3-2"
    key = "generated/mwa_abrechnung_3_2/parquet/mwa_abrechnung_entities/part-00001.parquet"

    def __init__(self) -> None:
        self.head_bucket_requests: list[str] = []
        self.list_object_buckets: list[str] = []

    def head_bucket(self, **kwargs):
        bucket = kwargs["Bucket"]
        self.head_bucket_requests.append(bucket)
        if bucket != self.bucket:
            raise AssertionError(f"Unexpected startup seed bucket probe: {bucket}")
        return {}

    def list_objects_v2(self, **kwargs):
        bucket = kwargs["Bucket"]
        self.list_object_buckets.append(bucket)
        if bucket != self.bucket:
            raise AssertionError(f"Unexpected startup seed bucket listing: {bucket}")
        return {
            "Contents": [{"Key": self.key}],
            "IsTruncated": False,
        }

    def head_object(self, **kwargs):
        if kwargs["Bucket"] != self.bucket or kwargs["Key"] != self.key:
            raise AssertionError(f"Unexpected startup seed object head request: {kwargs}")
        return {
            "ETag": '"seed-etag"',
            "ContentLength": 512,
            "Metadata": {},
        }


class CsvS3DiscoveryTests(TestCase):
    def test_drop_discovered_relation_object_handles_stale_table(self) -> None:
        connection = duckdb.connect(":memory:")
        try:
            connection.execute("CREATE SCHEMA stale")
            connection.execute("CREATE TABLE stale.cached_relation AS SELECT 1 AS value")

            drop_discovered_relation_object(
                connection,
                schema_name="stale",
                relation_name="cached_relation",
            )

            exists = connection.execute(
                """
                SELECT COUNT(*)
                FROM information_schema.tables
                WHERE table_schema = 'stale'
                  AND table_name = 'cached_relation'
                """
            ).fetchone()[0]
            self.assertEqual(exists, 0)
        finally:
            connection.close()

    def test_s3_existing_specs_ignore_loader_cache_tables(self) -> None:
        connection = duckdb.connect(":memory:")
        try:
            connection.execute("CREATE SCHEMA generated")
            connection.execute("CREATE VIEW generated.s3_relation AS SELECT 1 AS value")
            connection.execute("CREATE TABLE generated.s3_relation_duckdb_cache AS SELECT 1 AS value")

            discoverer = S3DataSourceDiscoverer(make_settings())
            specs = discoverer._load_existing_specs(connection)

            self.assertIn("generated.s3_relation", specs)
            self.assertNotIn("generated.s3_relation_duckdb_cache", specs)
        finally:
            connection.close()

    def test_discovered_csv_spec_uses_uploaded_csv_metadata_for_query_sql(self) -> None:
        discoverer = S3DataSourceDiscoverer(make_settings())

        specs = discoverer._build_desired_specs(FakeS3Client(), {"vat-smoke-test"})
        self.assertEqual(len(specs), 1)
        spec = next(iter(specs.values()))

        self.assertEqual(spec.relation_name, "tax_office")
        self.assertEqual(spec.object_revision, "abc123|128")
        self.assertEqual(spec.csv_delimiter, ",")
        self.assertTrue(spec.csv_has_header)
        self.assertIn("HEADER = TRUE", spec.query_sql)
        self.assertIn("DELIM = ','", spec.query_sql)

    def test_s3_discovery_ignores_data_exchange_prefix_objects(self) -> None:
        discoverer = S3DataSourceDiscoverer(make_settings())

        specs = discoverer._build_desired_specs(DataExchangeHiddenS3Client(), {"vat-smoke-test"})

        self.assertEqual(len(specs), 1)
        spec = next(iter(specs.values()))
        self.assertEqual(spec.display_name, "visible.csv")
        self.assertNotIn("data-exchange", spec.object_path)

    def test_partitioned_parquet_upload_is_discovered_as_one_relation(self) -> None:
        discoverer = S3DataSourceDiscoverer(make_settings())

        specs = discoverer._build_desired_specs(PartitionedParquetS3Client(), {"vat-smoke-test"})

        self.assertEqual(len(specs), 1)
        spec = next(iter(specs.values()))
        self.assertEqual(spec.relation_name, "vat_smoke")
        self.assertEqual(spec.object_path, "s3://vat-smoke-test/incoming/april/vat_smoke/**/*.parquet")
        self.assertIn("read_parquet(", spec.query_sql)
        self.assertIn("hive_partitioning=true", spec.query_sql)
        self.assertEqual(spec.display_name, "vat_smoke.parquet")
        self.assertEqual(spec.download_kind, "partitioned_parts")
        self.assertEqual(spec.part_prefix, "incoming/april/vat_smoke/")
        self.assertEqual(spec.part_count, 2)
        self.assertTrue(spec.zip_downloadable)
        service = object.__new__(WorkbenchService)
        service._data_source_discovery = SimpleNamespace(
            s3_relation_specs=lambda: specs,
        )
        metadata = WorkbenchService._s3_storage_object_metadata(service)
        item = next(iter(metadata.values()))
        self.assertEqual(
            item["query_alias"],
            "s3.vat_smoke_test.incoming.april.vat_smoke.parquet",
        )
        self.assertNotIn("vat_smoke.vat_smoke", item["query_alias"])

    def test_bounded_s3_sync_only_scans_requested_startup_seed_buckets(self) -> None:
        client = StartupSeedBucketS3Client()
        with TemporaryDirectory() as temp_dir:
            connection = duckdb.connect(":memory:")
            try:
                connection.execute("CREATE SCHEMA other_bucket")
                connection.execute("CREATE VIEW other_bucket.keep_me AS SELECT 1 AS value")
                discoverer = S3DataSourceDiscoverer(
                    make_temp_settings(Path(temp_dir) / "workspace.duckdb")
                )

                with (
                    patch(
                        "bit_data_workbench.backend.source_discovery.s3_client",
                        return_value=client,
                    ),
                    patch(
                        "bit_data_workbench.backend.source_discovery.list_s3_buckets",
                        side_effect=AssertionError("bounded startup discovery must not list every S3 bucket"),
                    ),
                    patch(
                        "bit_data_workbench.backend.source_discovery.build_s3_query",
                        return_value="SELECT 1 AS value",
                    ),
                ):
                    result = discoverer.sync_buckets(connection, [client.bucket])

                self.assertIsNotNone(result)
                self.assertEqual(client.head_bucket_requests, [client.bucket])
                self.assertEqual(client.list_object_buckets, [client.bucket])
                seed_schema = s3_bucket_schema_name(client.bucket)
                discovered = {
                    tuple(row)
                    for row in connection.execute(
                        """
                        SELECT table_schema, table_name
                        FROM information_schema.tables
                        WHERE table_schema IN (?, 'other_bucket')
                        ORDER BY table_schema, table_name
                        """,
                        [seed_schema],
                    ).fetchall()
                }
                self.assertIn(("other_bucket", "keep_me"), discovered)
                self.assertIn((seed_schema, "mwa_abrechnung_entities_parquet"), discovered)
            finally:
                connection.close()

    def test_startup_seed_source_sync_uses_bounded_s3_buckets(self) -> None:
        calls: list[tuple[tuple[str, ...], bool]] = []
        service = object.__new__(WorkbenchService)
        service._log_startup = lambda *_args, **_kwargs: None
        service._data_source_discovery = SimpleNamespace(
            sync_s3_buckets=lambda buckets, *, emit_event=True: calls.append((tuple(buckets), emit_event)),
            sync_source=lambda *_args, **_kwargs: (_ for _ in ()).throw(
                AssertionError("startup must not run full S3 source discovery")
            ),
        )

        WorkbenchService._sync_startup_seed_data_sources(service)

        self.assertEqual(calls, [(STARTUP_SEED_S3_BUCKETS, False)])

    def test_build_s3_parquet_query_accepts_runtime_hive_option(self) -> None:
        self.assertEqual(
            build_s3_query(
                "parquet",
                "s3://vat-smoke-test/incoming/april/vat_smoke/**/*.parquet",
                hive_partitioning=True,
            ),
            "SELECT * FROM read_parquet('s3://vat-smoke-test/incoming/april/vat_smoke/**/*.parquet', hive_partitioning=true)",
        )
        self.assertIn(
            "hive_partitioning=false",
            build_s3_query(
                "parquet",
                "s3://vat-smoke-test/incoming/april/vat_smoke/**/*.parquet",
                hive_partitioning=False,
            ),
        )

    def test_discovered_xml_and_xlsx_specs_materialize_to_queryable_csv_views(self) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["record_id", "tax_office"])
        worksheet.append([1, "Zurich Central Tax Office"])
        output = BytesIO()
        workbook.save(output)
        workbook.close()

        class QueryableFileClient:
            def list_objects_v2(self, **kwargs):
                bucket = kwargs["Bucket"]
                if bucket != "vat-smoke-test":
                    return {"Contents": [], "IsTruncated": False}
                return {
                    "Contents": [
                        {"Key": "incoming/tax-office.xml"},
                        {"Key": "incoming/tax-office.xlsx"},
                    ],
                    "IsTruncated": False,
                }

            def head_object(self, **kwargs):
                key = kwargs["Key"]
                if key.endswith(".xml"):
                    return {
                        "ETag": '"xml123"',
                        "ContentLength": 144,
                        "Metadata": {},
                    }
                if key.endswith(".xlsx"):
                    return {
                        "ETag": '"xlsx123"',
                        "ContentLength": 256,
                        "Metadata": {},
                    }
                raise AssertionError("Unexpected head_object request.")

            def get_object(self, **kwargs):
                key = kwargs["Key"]
                if key.endswith(".xml"):
                    return {
                        "Body": BytesIO(
                            b"<rows><row><record_id>1</record_id><tax_office>Zurich Central Tax Office</tax_office></row></rows>"
                        )
                    }
                if key.endswith(".xlsx"):
                    return {"Body": BytesIO(output.getvalue())}
                raise AssertionError("Unexpected get_object request.")

        with TemporaryDirectory() as temp_dir:
            discoverer = S3DataSourceDiscoverer(
                make_temp_settings(Path(temp_dir) / "workspace.duckdb")
            )

            specs = discoverer._build_desired_specs(QueryableFileClient(), {"vat-smoke-test"})

            xml_spec = next(spec for spec in specs.values() if spec.display_name == "tax-office.xml")
            self.assertEqual(xml_spec.object_format, "xml")
            self.assertEqual(xml_spec.display_name, "tax-office.xml")
            self.assertEqual(xml_spec.size_bytes, 144)
            self.assertIn("read_csv_auto(", xml_spec.query_sql)
            self.assertTrue((Path(temp_dir) / "s3-query-sources").exists())

            xlsx_spec = next(spec for spec in specs.values() if spec.display_name == "tax-office.xlsx")
            self.assertEqual(xlsx_spec.object_format, "xlsx")
            self.assertEqual(xlsx_spec.display_name, "tax-office.xlsx")
            self.assertEqual(xlsx_spec.size_bytes, 256)
            self.assertIn("read_csv_auto(", xlsx_spec.query_sql)

    def test_generated_loader_multiformat_outputs_use_file_display_names(self) -> None:
        discoverer = S3DataSourceDiscoverer(make_settings())

        specs = discoverer._build_desired_specs(
            GeneratedMwaLoaderS3Client(),
            {"vat-smoke-test"},
        )
        specs_by_relation = {spec.relation_name: spec for spec in specs.values()}

        self.assertNotIn("mwa_abrechnung_test", specs_by_relation)

        csv_spec = specs_by_relation["mwa_abrechnung_entities_csv"]
        self.assertEqual(csv_spec.display_name, "mwa_abrechnung_entities.csv")
        self.assertEqual(csv_spec.object_format, "csv")
        self.assertTrue(csv_spec.object_path.endswith("/csv/mwa_abrechnung_entities/*.csv"))
        self.assertEqual(csv_spec.size_bytes, 224)
        self.assertTrue(csv_spec.csv_has_header)
        self.assertEqual(csv_spec.csv_delimiter, ",")
        self.assertEqual(csv_spec.download_kind, "generated_parts")
        self.assertEqual(
            csv_spec.part_prefix,
            "generated/mwa_abrechnung_test/csv/mwa_abrechnung_entities/",
        )
        self.assertEqual(csv_spec.part_file_format, "csv")
        self.assertEqual(csv_spec.part_count, 2)
        self.assertEqual(csv_spec.download_filename, "mwa_abrechnung_entities.csv")
        self.assertTrue(csv_spec.merge_downloadable)
        self.assertTrue(csv_spec.zip_downloadable)
        self.assertIn("read_csv_auto(", csv_spec.query_sql)
        self.assertIn("HEADER = TRUE", csv_spec.query_sql)

        json_spec = specs_by_relation["mwa_abrechnung_entities_json"]
        self.assertEqual(json_spec.display_name, "mwa_abrechnung_entities.jsonl")
        self.assertEqual(json_spec.object_format, "jsonl")
        self.assertTrue(json_spec.object_path.endswith("/json/mwa_abrechnung_entities/*.jsonl"))
        self.assertEqual(json_spec.size_bytes, 448)
        self.assertEqual(json_spec.download_kind, "generated_parts")
        self.assertEqual(
            json_spec.part_prefix,
            "generated/mwa_abrechnung_test/json/mwa_abrechnung_entities/",
        )
        self.assertEqual(json_spec.part_file_format, "jsonl")
        self.assertEqual(json_spec.part_count, 2)
        self.assertEqual(json_spec.download_filename, "mwa_abrechnung_entities.jsonl")
        self.assertTrue(json_spec.merge_downloadable)
        self.assertTrue(json_spec.zip_downloadable)
        self.assertIn("read_json_auto(", json_spec.query_sql)

        parquet_spec = specs_by_relation["mwa_abrechnung_entities_parquet"]
        self.assertEqual(parquet_spec.display_name, "mwa_abrechnung_entities.parquet")
        self.assertEqual(parquet_spec.object_format, "parquet")
        self.assertTrue(parquet_spec.object_path.endswith("/parquet/mwa_abrechnung_entities/*.parquet"))
        self.assertEqual(parquet_spec.size_bytes, 896)
        self.assertEqual(parquet_spec.download_kind, "generated_parts")
        self.assertEqual(
            parquet_spec.part_prefix,
            "generated/mwa_abrechnung_test/parquet/mwa_abrechnung_entities/",
        )
        self.assertEqual(parquet_spec.part_file_format, "parquet")
        self.assertEqual(parquet_spec.part_count, 2)
        self.assertEqual(parquet_spec.download_filename, "mwa_abrechnung_entities.parquet")
        self.assertFalse(parquet_spec.merge_downloadable)
        self.assertTrue(parquet_spec.zip_downloadable)

    def test_single_part_generated_loader_output_keeps_dataset_wildcard_reference(self) -> None:
        discoverer = S3DataSourceDiscoverer(make_settings())

        specs = discoverer._build_desired_specs(
            SinglePartGeneratedMwaLoaderS3Client(),
            {"vat-smoke-test"},
        )
        specs_by_relation = {spec.relation_name: spec for spec in specs.values()}

        parquet_spec = specs_by_relation["mwa_abrechnung_entities_parquet"]
        self.assertEqual(parquet_spec.display_name, "mwa_abrechnung_entities.parquet")
        self.assertEqual(parquet_spec.object_format, "parquet")
        self.assertTrue(
            parquet_spec.object_path.endswith(
                "/parquet/mwa_abrechnung_entities/*.parquet"
            )
        )
        self.assertNotIn("part-00001.parquet", parquet_spec.object_path)
        self.assertEqual(parquet_spec.download_kind, "generated_parts")
        self.assertEqual(parquet_spec.part_count, 1)
        self.assertEqual(
            parquet_spec.part_prefix,
            "generated/mwa_abrechnung_test/parquet/mwa_abrechnung_entities/",
        )

    def test_generated_parts_expose_download_metadata_in_workspace_metadata(self) -> None:
        discoverer = S3DataSourceDiscoverer(make_settings())
        specs = discoverer._build_desired_specs(
            GeneratedMwaLoaderS3Client(),
            {"vat-smoke-test"},
        )
        service = object.__new__(WorkbenchService)
        service._data_source_discovery = SimpleNamespace(
            s3_relation_specs=lambda: specs,
        )

        metadata = WorkbenchService._s3_storage_object_metadata(service)
        by_display_name = {
            value["display_name"]: value
            for value in metadata.values()
        }

        csv_metadata = by_display_name["mwa_abrechnung_entities.csv"]
        self.assertFalse(csv_metadata["downloadable"])
        self.assertEqual(csv_metadata["download_kind"], "generated_parts")
        self.assertEqual(csv_metadata["part_count"], 2)
        self.assertEqual(csv_metadata["size_bytes"], 224)
        self.assertEqual(csv_metadata["download_filename"], "mwa_abrechnung_entities.csv")
        self.assertTrue(csv_metadata["merge_downloadable"])
        self.assertTrue(csv_metadata["zip_downloadable"])
        self.assertEqual(
            csv_metadata["query_alias"],
            "s3.vat_smoke_test.generated.mwa_abrechnung_test.csv.mwa_abrechnung_entities.csv",
        )
        self.assertNotIn("mwa_abrechnung_entities.mwa_abrechnung_entities", csv_metadata["query_alias"])

        json_metadata = by_display_name["mwa_abrechnung_entities.jsonl"]
        self.assertFalse(json_metadata["downloadable"])
        self.assertEqual(json_metadata["download_kind"], "generated_parts")
        self.assertEqual(json_metadata["part_count"], 2)
        self.assertEqual(json_metadata["size_bytes"], 448)
        self.assertEqual(json_metadata["download_filename"], "mwa_abrechnung_entities.jsonl")
        self.assertTrue(json_metadata["merge_downloadable"])
        self.assertTrue(json_metadata["zip_downloadable"])
        self.assertEqual(
            json_metadata["query_alias"],
            "s3.vat_smoke_test.generated.mwa_abrechnung_test.json.mwa_abrechnung_entities.jsonl",
        )
        self.assertNotIn("mwa_abrechnung_entities.mwa_abrechnung_entities", json_metadata["query_alias"])

        parquet_metadata = by_display_name["mwa_abrechnung_entities.parquet"]
        self.assertFalse(parquet_metadata["downloadable"])
        self.assertEqual(parquet_metadata["download_kind"], "generated_parts")
        self.assertEqual(parquet_metadata["part_count"], 2)
        self.assertEqual(parquet_metadata["size_bytes"], 896)
        self.assertEqual(parquet_metadata["download_filename"], "mwa_abrechnung_entities.parquet")
        self.assertFalse(parquet_metadata["merge_downloadable"])
        self.assertTrue(parquet_metadata["zip_downloadable"])
        self.assertEqual(
            parquet_metadata["query_alias"],
            "s3.vat_smoke_test.generated.mwa_abrechnung_test.parquet.mwa_abrechnung_entities.parquet",
        )
        self.assertNotIn("mwa_abrechnung_entities.mwa_abrechnung_entities", parquet_metadata["query_alias"])

    def test_generated_parts_workspace_metadata_normalizes_stale_single_part_path(self) -> None:
        service = object.__new__(WorkbenchService)
        service._data_source_discovery = SimpleNamespace(
            s3_relation_specs=lambda: {
                "kostenbelege.dim_kalender_parquet": DiscoveredRelationSpec(
                    schema_name="kostenbelege",
                    relation_name="dim_kalender_parquet",
                    query_sql=(
                        "SELECT * FROM read_parquet("
                        "'s3://poc-tests-performance-evaluation-kostenbelege-3-1/"
                        "generated/kostenbelege_3_1/parquet/dim_kalender/part-00001.parquet')"
                    ),
                    object_path=(
                        "s3://poc-tests-performance-evaluation-kostenbelege-3-1/"
                        "generated/kostenbelege_3_1/parquet/dim_kalender/part-00001.parquet"
                    ),
                    object_format="parquet",
                    display_name="dim_kalender.parquet",
                    download_kind="generated_parts",
                    part_prefix="generated/kostenbelege_3_1/parquet/dim_kalender/",
                    part_file_format="parquet",
                    part_count=1,
                    download_filename="dim_kalender.parquet",
                    zip_downloadable=True,
                )
            },
        )

        metadata = WorkbenchService._s3_storage_object_metadata(service)
        item = metadata["kostenbelege.dim_kalender_parquet"]

        self.assertEqual(
            item["path"],
            "s3://poc-tests-performance-evaluation-kostenbelege-3-1/"
            "generated/kostenbelege_3_1/parquet/dim_kalender/*.parquet",
        )
        self.assertEqual(
            item["query_reference"],
            's3."poc-tests-performance-evaluation-kostenbelege-3-1".'
            '"generated/kostenbelege_3_1/parquet/dim_kalender/*.parquet"',
        )
        self.assertIn("dim_kalender/*.parquet", item["query_sql"])
        self.assertNotIn("part-00001.parquet", item["query_reference"])
        self.assertNotIn("part-00001.parquet", item["query_sql"])
        self.assertEqual(item["download_kind"], "generated_parts")
        self.assertEqual(item["key"], "")

    def test_cached_generated_source_object_normalizes_stale_single_part_path(self) -> None:
        source_object = SourceObject(
            name="dim_kalender_parquet",
            kind="PARQUET",
            relation="kostenbelege.dim_kalender_parquet",
            display_name="dim_kalender.parquet",
            query_reference=(
                's3."poc-tests-performance-evaluation-kostenbelege-3-1".'
                '"generated/kostenbelege_3_1/parquet/dim_kalender/part-00001.parquet"'
            ),
            query_sql=(
                "read_parquet('s3://poc-tests-performance-evaluation-kostenbelege-3-1/"
                "generated/kostenbelege_3_1/parquet/dim_kalender/part-00001.parquet')"
            ),
            s3_bucket="poc-tests-performance-evaluation-kostenbelege-3-1",
            s3_key="",
            s3_path=(
                "s3://poc-tests-performance-evaluation-kostenbelege-3-1/"
                "generated/kostenbelege_3_1/parquet/dim_kalender/part-00001.parquet"
            ),
            s3_file_format="parquet",
            s3_download_kind="generated_parts",
            s3_part_prefix="generated/kostenbelege_3_1/parquet/dim_kalender/",
            s3_part_file_format="parquet",
            s3_part_count=1,
            s3_download_filename="dim_kalender.parquet",
            s3_zip_downloadable=True,
        )

        WorkbenchService._normalize_generated_s3_source_object(source_object)

        self.assertEqual(
            source_object.s3_path,
            "s3://poc-tests-performance-evaluation-kostenbelege-3-1/"
            "generated/kostenbelege_3_1/parquet/dim_kalender/*.parquet",
        )
        self.assertEqual(
            source_object.query_reference,
            's3."poc-tests-performance-evaluation-kostenbelege-3-1".'
            '"generated/kostenbelege_3_1/parquet/dim_kalender/*.parquet"',
        )
        self.assertEqual(source_object.s3_key, "")
        self.assertIn("dim_kalender/*.parquet", source_object.query_sql)
        self.assertNotIn("part-00001.parquet", source_object.query_reference)
        self.assertNotIn("part-00001.parquet", source_object.query_sql)
